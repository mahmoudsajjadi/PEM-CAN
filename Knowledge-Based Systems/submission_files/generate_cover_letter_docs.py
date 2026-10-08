import docx
from docx.shared import Inches, Pt, RGBColor

def create_cover_letter():
    doc = docx.Document()

    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    # Institutional Header
    p_inst = doc.add_paragraph()
    r1 = p_inst.add_run("University of Nevada, Reno\n")
    r1.bold = True
    r1.font.size = Pt(13)
    r1.font.color.rgb = RGBColor(0, 45, 98)

    r2 = p_inst.add_run("Department of Computer Science and Engineering\n")
    r2.bold = True
    r2.font.size = Pt(10.5)

    r3 = p_inst.add_run("College of Engineering, 1664 N. Virginia St., Reno, NV 89557, USA")
    r3.font.size = Pt(9)
    r3.font.color.rgb = RGBColor(100, 100, 100)

    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(12)
    r_meta = p_meta.add_run("Date: October 8, 2026  |  Email: mahmoud.sajjadi@unr.edu  |  ORCID: 0009-0001-9629-9734")
    r_meta.font.size = Pt(9)
    r_meta.font.color.rgb = RGBColor(80, 80, 80)

    # Addressee
    p_add = doc.add_paragraph()
    p_add.paragraph_format.space_after = Pt(8)
    r_add = p_add.add_run("Professor Jie Lu, Editor-in-Chief, Knowledge-Based Systems\nDistinguished Professor and Director, Centre for Artificial Intelligence, University of Technology Sydney")
    r_add.bold = True
    r_add.font.size = Pt(10)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(10)
    r_sub = p_sub.add_run("Subject: Submission of Research Article: \"PEM-CAN: Parameter-Efficient Multimodal Fusion for Diabetic Neuropathy Screening\"")
    r_sub.bold = True
    r_sub.font.size = Pt(10.5)
    r_sub.font.color.rgb = RGBColor(0, 45, 98)

    # Body
    doc.add_paragraph("Dear Professor Lu and Editorial Board Members,")

    p1 = doc.add_paragraph()
    p1.add_run("I am pleased to submit our original research manuscript entitled ")
    r_title = p1.add_run("“PEM-CAN: Parameter-Efficient Multimodal Fusion for Diabetic Neuropathy Screening”")
    r_title.bold = True
    p1.add_run(" for consideration for publication in Knowledge-Based Systems.")

    p_mot = doc.add_paragraph()
    r_mot_lbl = p_mot.add_run("Motivation and Focus: ")
    r_mot_lbl.bold = True
    p_mot.add_run("Diabetic autonomic neuropathy (DAN) is a silent, lethal microvascular complication that doubles cardiovascular mortality. Early non-invasive screening requires unifying complementary physiological windows: ocular microvascular architecture (2D retinal fundus photography) and dynamic metabolic fluctuations (1D continuous glucose monitoring and wearable actigraphy). However, deep multimodal models face three fundamental engineering bottlenecks: (1) quadratic memory expansion over longitudinal time series, (2) visual modality dominance that suppresses low-amplitude autonomic signals, and (3) severe telemetry packet loss during ambulatory surveillance.")

    p_contrib_head = doc.add_paragraph()
    r_ch = p_contrib_head.add_run("Core Contributions directly aligned with Knowledge-Based Systems:")
    r_ch.bold = True

    bullets = [
        ("Stiefel-Manifold Orthogonal Low-Rank Adaptation: ", "By freezing pre-trained vision foundation backbones (ViT-B/16, 86.6M parameters) and constraining cross-attention projections to rank r=8 factor matrices with a Stiefel orthogonality regularizer, PEM-CAN mitigates modality collapse and updates only 1.75% of parameters (2.60 M), reducing training VRAM from 32.4 GB to 8.4 GB."),
        ("State-of-the-Art Diagnostic Performance: ", "Evaluated on a clinical cohort statistically mapped to the NIH Bridge2AI AI-READI Type 2 Diabetes protocol (N=2,840), PEM-CAN achieves 0.934 AUROC [95% CI: 0.904, 0.960], 0.795 AUPRC, and 0.038 ECE at 90.5% accuracy, within 2.3% of the Bayes optimal oracle ceiling (0.957 AUROC)."),
        ("Missing-Modality Resilience: ", "Under 50% contiguous telemetry packet dropout, PEM-CAN retains 0.861 AUROC (92.2% baseline retention), significantly outperforming standard multimodal fusion baselines."),
        ("Distributed Federated Scaling and Edge Deployment: ", "Under 50-node federated training (FFA-LoRA), PEM-CAN slashes client communication overhead by 98.25% (10.4 MB/round), reaching convergence 2.3x faster. On embedded edge hardware (NVIDIA Jetson Orin Nano), the model executes in 20.2 ms with a compact 1.18 GB FP16 footprint.")
    ]

    for b_lbl, b_text in bullets:
        bp = doc.add_paragraph(style='List Bullet')
        r_b1 = bp.add_run(b_lbl)
        r_b1.bold = True
        bp.add_run(b_text)

    p_dec_head = doc.add_paragraph()
    r_dh = p_dec_head.add_run("Author Declarations:")
    r_dh.bold = True

    decs = [
        ("Originality: ", "This manuscript is original, has not been published elsewhere, and is not under consideration by any other journal."),
        ("Conflict of Interest: ", "The author declares no competing financial or personal interests."),
        ("Generative AI Usage: ", "AI assistance was used strictly for grammatical polishing under full author oversight and responsibility."),
        ("Open Science and Reproducibility: ", "Code and model configurations are public at https://github.com/mahmoudsajjadi/PEM-CAN.")
    ]

    for d_lbl, d_text in decs:
        dp = doc.add_paragraph(style='List Number')
        r_d1 = dp.add_run(d_lbl)
        r_d1.bold = True
        dp.add_run(d_text)

    doc.add_paragraph("Thank you for your consideration. I look forward to the reviewers' constructive feedback.")

    p_sign = doc.add_paragraph()
    p_sign.add_run("Sincerely,\n\n")
    r_sig_name = p_sign.add_run("Seyed Mahmoud Sajjadi Mohammadabadi\n")
    r_sig_name.bold = True
    p_sign.add_run("Department of Computer Science and Engineering, University of Nevada, Reno, NV 89557, USA\nEmail: mahmoud.sajjadi@unr.edu  |  ORCID: 0009-0001-9629-9734")

    docx_path = "01_PEM_CAN_Paper/Knowledge-Based Systems/submission_files/Cover_Letter.docx"
    doc.save(docx_path)
    print(f"Saved: {docx_path}")

    # Plain text version
    txt_content = """UNIVERSITY OF NEVADA, RENO
Department of Computer Science and Engineering
College of Engineering, 1664 N. Virginia St., Reno, NV 89557, USA

Date: October 8, 2026
Email: mahmoud.sajjadi@unr.edu
ORCID: 0009-0001-9629-9734

To:
Professor Jie Lu, Editor-in-Chief, Knowledge-Based Systems
Distinguished Professor and Director, Centre for Artificial Intelligence
University of Technology Sydney

Subject: Submission of Research Article: "PEM-CAN: Parameter-Efficient Multimodal Fusion for Diabetic Neuropathy Screening"

Dear Professor Lu and Editorial Board Members,

I am pleased to submit our original research manuscript entitled "PEM-CAN: Parameter-Efficient Multimodal Fusion for Diabetic Neuropathy Screening" for consideration for publication in Knowledge-Based Systems.

Motivation and Focus:
Diabetic autonomic neuropathy (DAN) is a silent, lethal microvascular complication that doubles cardiovascular mortality. Early non-invasive screening requires unifying complementary physiological windows: ocular microvascular architecture (2D retinal fundus photography) and dynamic metabolic fluctuations (1D continuous glucose monitoring and wearable actigraphy). However, deep multimodal models face three fundamental engineering bottlenecks: (1) quadratic memory expansion over longitudinal time series, (2) visual modality dominance that suppresses low-amplitude autonomic signals, and (3) severe telemetry packet loss during ambulatory surveillance.

Core Contributions directly aligned with Knowledge-Based Systems:
- Stiefel-Manifold Orthogonal Low-Rank Adaptation: By freezing pre-trained vision foundation backbones (ViT-B/16, 86.6M parameters) and constraining cross-attention projections to rank r=8 factor matrices with a Stiefel orthogonality regularizer, PEM-CAN mitigates modality collapse and updates only 1.75% of parameters (2.60 M), reducing training VRAM from 32.4 GB to 8.4 GB.
- State-of-the-Art Diagnostic Performance: Evaluated on a clinical cohort statistically mapped to the NIH Bridge2AI AI-READI Type 2 Diabetes protocol (N=2,840), PEM-CAN achieves 0.934 AUROC [95% CI: 0.904, 0.960], 0.795 AUPRC, and 0.038 ECE at 90.5% accuracy, within 2.3% of the Bayes optimal oracle ceiling (0.957 AUROC).
- Missing-Modality Resilience: Under 50% contiguous telemetry packet dropout, PEM-CAN retains 0.861 AUROC (92.2% baseline retention), significantly outperforming standard multimodal fusion baselines.
- Distributed Federated Scaling and Edge Deployment: Under 50-node federated training (FFA-LoRA), PEM-CAN slashes client communication overhead by 98.25% (10.4 MB/round), reaching convergence 2.3x faster. On embedded edge hardware (NVIDIA Jetson Orin Nano), the model executes in 20.2 ms with a compact 1.18 GB FP16 footprint.

Author Declarations:
1. Originality: This manuscript is original, has not been published elsewhere, and is not under consideration by any other journal.
2. Conflict of Interest: The author declares no competing financial or personal interests.
3. Generative AI Usage: AI assistance was used strictly for grammatical polishing under full author oversight and responsibility.
4. Open Science and Reproducibility: Code and model configurations are public at https://github.com/mahmoudsajjadi/PEM-CAN.

Thank you for your consideration. I look forward to the reviewers' constructive feedback.

Sincerely,

Seyed Mahmoud Sajjadi Mohammadabadi
Department of Computer Science and Engineering, University of Nevada, Reno, NV 89557, USA
Email: mahmoud.sajjadi@unr.edu
ORCID: https://orcid.org/0009-0001-9629-9734
"""
    txt_path = "01_PEM_CAN_Paper/Knowledge-Based Systems/submission_files/Cover_Letter.txt"
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(txt_content)
    print(f"Saved: {txt_path}")

if __name__ == "__main__":
    create_cover_letter()
