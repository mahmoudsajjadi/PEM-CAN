import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT_DIR = r"C:\Users\17756\OneDrive - University of Nevada, Reno\Projects\Fast papers\Sustainability\Review\01_PEM_CAN_Paper\Knowledge-Based Systems\submission_files"
os.makedirs(OUT_DIR, exist_ok=True)

# ---------------------------------------------------------------------------
# 1. Highlights.docx
# ---------------------------------------------------------------------------
doc_hl = docx.Document()

# Adjust margins
for section in doc_hl.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

title_p = doc_hl.add_paragraph()
title_run = title_p.add_run("Highlights")
title_run.font.name = "Arial"
title_run.font.size = Pt(16)
title_run.font.bold = True
title_run.font.color.rgb = RGBColor(0, 45, 98)

sub_p = doc_hl.add_paragraph()
sub_run = sub_p.add_run("PEM-CAN: Parameter-Efficient Multimodal Fusion for Diabetic Neuropathy Screening")
sub_run.font.name = "Arial"
sub_run.font.size = Pt(11)
sub_run.font.italic = True
sub_run.font.color.rgb = RGBColor(100, 100, 100)

doc_hl.add_paragraph() # Spacing

highlights = [
    "PEM-CAN fuses retinal images with continuous sensor streams for neuropathy staging.",
    "Orthogonal Stiefel adapters update 1.75% weights, slashing VRAM from 32.4 to 8.4 GB.",
    "Model achieves 0.934 AUROC and retains 0.861 AUROC under 50% sensor packet loss.",
    "FFA-LoRA federated training slashes client communication bandwidth by 98.25%.",
    "Edge deployment achieves 20.2 ms forward latency on embedded Jetson Orin Nano."
]

for hl in highlights:
    p = doc_hl.add_paragraph(style='List Bullet')
    r = p.add_run(hl)
    r.font.name = "Arial"
    r.font.size = Pt(11)

hl_path = os.path.join(OUT_DIR, "Highlights.docx")
doc_hl.save(hl_path)
print(f"Saved: {hl_path}")

# Also save plain text Highlights.txt
txt_path = os.path.join(OUT_DIR, "Highlights.txt")
with open(txt_path, "w", encoding="utf-8") as f:
    f.write("Highlights\n")
    f.write("PEM-CAN: Parameter-Efficient Multimodal Fusion for Diabetic Neuropathy Screening\n\n")
    for hl in highlights:
        f.write(f"- {hl} ({len(hl)} chars)\n")
print(f"Saved: {txt_path}")

# ---------------------------------------------------------------------------
# 2. Declaration_of_Competing_Interests.docx
# ---------------------------------------------------------------------------
doc_dec = docx.Document()

for section in doc_dec.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

h_p = doc_dec.add_paragraph()
h_run = h_p.add_run("Declaration of Competing Interests")
h_run.font.name = "Arial"
h_run.font.size = Pt(16)
h_run.font.bold = True
h_run.font.color.rgb = RGBColor(0, 45, 98)

p_meta = doc_dec.add_paragraph()
r_meta = p_meta.add_run("Manuscript Title: PEM-CAN: Parameter-Efficient Multimodal Fusion for Diabetic Neuropathy Screening\nJournal: Knowledge-Based Systems (Elsevier)\nAuthor: Seyed Mahmoud Sajjadi Mohammadabadi\nAffiliation: Department of Computer Science and Engineering, University of Nevada, Reno, NV 89557, USA")
r_meta.font.name = "Arial"
r_meta.font.size = Pt(10)
r_meta.font.color.rgb = RGBColor(80, 80, 80)

doc_dec.add_paragraph() # Spacing

p_body = doc_dec.add_paragraph()
r_body = p_body.add_run(
    "Declaration:\n\n"
    "The author declares that he has no known competing financial interests or personal relationships "
    "that could have appeared to influence the work reported in this paper.\n\n"
    "Specifically, the author confirms:\n"
    "1. No financial support or funding has been received that influenced the study outcome, analysis, or manuscript preparation.\n"
    "2. No employment, consultancies, honoraria, stock ownership, or options exist that relate to the subject matter of the manuscript.\n"
    "3. No patents, royalties, or licensing agreements are held or pending that conflict with this publication.\n"
    "4. No personal, professional, or academic relationships exist that could inappropriately influence or bias the research presented.\n"
    "5. The author does not currently serve, nor has previously served, in an editorial capacity for Knowledge-Based Systems."
)
r_body.font.name = "Arial"
r_body.font.size = Pt(11)

p_sign = doc_dec.add_paragraph()
p_sign.paragraph_format.space_before = Pt(20)
r_sign = p_sign.add_run(
    "Confirmed and certified by:\n\n"
    "Seyed Mahmoud Sajjadi Mohammadabadi\n"
    "Department of Computer Science and Engineering\n"
    "University of Nevada, Reno\n"
    "Date: October 8, 2026"
)
r_sign.font.name = "Arial"
r_sign.font.size = Pt(10.5)
r_sign.font.bold = True

dec_path = os.path.join(OUT_DIR, "Declaration_of_Competing_Interests.docx")
doc_dec.save(dec_path)
print(f"Saved: {dec_path}")
